"""Package the read-only analysis and application ledger into a new review workbook."""
import json, collections, hashlib, csv
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table,TableStyleInfo
ROOT=Path(__file__).resolve().parents[1]
read=lambda p:json.loads(p.read_text(encoding='utf8'))
rows=read(ROOT/'reference-assets/iv-master-review-2026-10-04/analysis.json')
ledger=read(ROOT/'docs/iv-master-aprovacoes-aplicadas-2026-10-04.json')
wb=Workbook();ws=wb.active;ws.title='Resumo'
for row in [
 ['Revisão IV Master — 04/10/2026','Resultado'],
 ['Alterações aprovadas','29 combinações em 21 registros existentes; 16 novas regras (64 combinações + 16 convenções 67%). Total 581 registros.'],
 ['Preservação','A planilha original com respostas não foi alterada. Pendências continuam como propostas; sem alterações em stats/famílias/resolvedor.'],
 ['Critério','Nível máximo 40; maior produto Ataque × Defesa × PS, com PS inteiro, dentro da mesma soma de IVs: 44/43/42/41. 67% = 10/10/10 intencional.'],
 ['Como validar','Use a aba Faltantes ou Formas e observações. Todos os IVs são texto para impedir conversão em data pelo Excel.'],
 ['Espécies faltantes','30 linhas F-017–F-046 (29 alvos distintos; Sinistcha Comum aparece duas vezes) + 4 registros N. Formas herdadas também podem precisar de regras específicas.'],
 ['Fontes','Stadium: atributos e fórmula do bundle atual; Dittobase: API atual e fórmula pública. As colunas são cálculos reproduzidos, não uma tabela oficial chamada IV Master por percentual.'],
 ['Empates','Listados todos os IVs de maior produto. A coluna Dittobase preserva empates após arredondamento; a interface pode ordenar por ataque/PS/PC.'],
 ['Poke Genie','Site oficial e descrição oficial consultados; sem dataset/API público completo encontrado. Não há confirmação numérica do Poke Genie nesta revisão.'],
 ['Cache do jogo','Extração somente leitura do aparelho 2412DPC0AG, em 04/10/2026: 2.480 templates. 2.477 compartilhados com 27/09 sem mudanças, 3 novos. Cache local não garante atualização integral do servidor.'],
 ['Goodra Hisui','Jogo/app: 211/255/190; Dittobase: 220/242/207. Priorizar propostas calculadas com o jogo.'],
 ['Basculegion','Macho e fêmea sem template no cache consultado. Propostas provisórias pelos stats locais, coincidentes com Dittobase; não confirmadas pelo jogo.'],
 ['Zygarde','A regra atual coincide com COMPLETE, forma 100%. Mantida. COMPLETE_FIFTY_PERCENT e COMPLETE_TEN_PERCENT têm stats das formas 50% e 10%, respectivamente.'],
 ['Famílias de evolução','Mantidas as regras GALARIANFARFETCHD e GLIMMET: coincidem com Sirfetch’d e Glimmora, conforme observações do usuário.'],
 ['Limites','Sem afirmar elegibilidade em copas ou disponibilidade de captura. Regra familiar pode servir à evolução; forma com mesmos IVs não implica mesmos stats.'],
 ['Fonte Stadium','https://www.stadiumgaming.gg/rank-checker'],
 ['Fonte Dittobase','https://www.dittobase.com/pokemon-go/pvp-ivs'],
 ['Fonte Poke Genie','https://pokegenie.net/'],
 ['Planilha original SHA-256',ledger['workbook_sha256']],
]:ws.append(row)

ws=wb.create_sheet('Aplicadas');ws.append(['ID revisão','Registro','Percentual','Antes','Depois','Observação'])
for r in ledger['changes']:ws.append([r['id'],r['key'],r['percent'],'/'.join(map(str,r['before'])) if r['before'] else 'Novo','/'.join(map(str,r['after'])),r['note']])
common=['ID','Espécie/forma','Alvo da regra','Situação','Atributos usados','Fonte dos atributos','Regra comparada','Atual 98','Atual 96','Atual 93','Atual 91','Proposta 98','Proposta 96','Proposta 93','Proposta 91','Observação','Decisão','IV 98 aprovado','IV 96 aprovado','IV 93 aprovado','IV 91 aprovado']
for title,predicate in [('Faltantes',lambda r:(r['id'].startswith('F-') and r['decisão original']!='Aprovar proposta') or r['id'].startswith('N-')),('Formas e observações',lambda r:r['id'].startswith('H-') or r['id']=='ANÁLISE')]:
 ws=wb.create_sheet(title);ws.append(common)
 for r in rows:
  if predicate(r):ws.append([r['id'],r['espécie/forma'],r['alvo'],r['situação'],r['atributos usados'],r['fonte atributos'],r['regra comparada'],*[r['atual '+p] for p in ('98','96','93','91')],*[r['proposta '+p] for p in ('98','96','93','91')],r['observação'],'Pendente','','','',''])
 dv=DataValidation(type='list',formula1='"Pendente,Aprovar proposta,Manter atual,Definir outro IV,Revisar forma/família"');ws.add_data_validation(dv);dv.add(f'Q2:Q{ws.max_row}')
ws=wb.create_sheet('Comparação fontes');ws.append(list(rows[0]))
for r in rows:ws.append(list(r.values()))
ws=wb.create_sheet('Respostas originais');ws.append(['Aba','ID','Registro/forma','Decisão','Resposta/observações','Tratamento'])
original=load_workbook(ROOT/'docs/revisao-iv-master-nivel-40.xlsx')
for sheet in list(original)[1:]:
 values=list(sheet.values);h=values[0];decision=h.index('Decisão');note=h.index('Observações do usuário')
 for r in values[1:]:
  result='Aplicado' if r[decision]=='Aprovar proposta' else 'Analisado; ver propostas e relatório'
  if str(r[0]).startswith('E-'):result='Mantido: já usa IVs da evolução solicitada'
  if str(r[0]) in ['D-030','D-031','D-032']:result='Mantido: confirmado Zygarde Completo/100%'
  ws.append([sheet.title,r[0],r[1],r[decision],r[note],result])
for idx,ws in enumerate(wb):
 ws.freeze_panes='C2';ws.sheet_view.showGridLines=False
 for c in ws[1]:c.fill=PatternFill('solid',fgColor='243B53');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True)
 for col in ws.columns:
  label=str(col[0].value);width=20
  if any(x in label.lower() for x in ['observa','tratamento','resposta','fonte','situação']):width=48
  if label in ('Registro','Espécie/forma','Alvo da regra'):width=27
  ws.column_dimensions[col[0].column_letter].width=width
  for c in col[1:]:
   c.alignment=Alignment(wrap_text=True,vertical='top')
   if any(x in label.lower() for x in ('proposta','aprovado','atual','antes','depois','stadium 9','dittobase 9')):c.number_format='@'
   if 'aprovado' in label.lower() or label=='Decisão':c.fill=PatternFill('solid',fgColor='FFF2CC')
   if isinstance(c.value,str) and c.value.startswith('https://'):c.hyperlink=c.value;c.font=Font(color='0563C1',underline='single')
 if idx==0:ws.column_dimensions['A'].width=32;ws.column_dimensions['B'].width=130
 else:
  tab=Table(displayName=f'Resultado{idx}',ref=ws.dimensions);tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(tab)
 for i in range(2,ws.max_row+1):ws.row_dimensions[i].height=60
path=ROOT/'docs/revisao-iv-master-resultado-2026-10-04.xlsx';wb.save(path)
check=load_workbook(path);assert check['Aplicadas'].max_row==110
assert hashlib.sha256((ROOT/'docs/revisao-iv-master-nivel-40.xlsx').read_bytes()).hexdigest()==ledger['workbook_sha256']
print(path);print([(s.title,s.max_row-1) for s in check])
