"""Gera uma planilha de revisão, sem modificar os assets do aplicativo."""
import json, re, unicodedata, math
from pathlib import Path
from urllib.parse import quote
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
def read(path): return json.loads((ROOT/path).read_text(encoding='utf-8'))
def raw(s): return re.sub('[^A-Z0-9]', '', unicodedata.normalize('NFD',s.replace('♀','F').replace('♂','M')).encode('ascii','ignore').decode().upper())
def norm(s):
    s=raw(s)
    for a,b in [('ALOLAN','ALOLA'),('GALARIAN','GALAR'),('HISUIAN','HISUI'),('PALDEAN','PALDEA')]:
        if s.startswith(a): s=s[len(a):]+b; break
    if s in ['LYCANROCDIURNO','LYCANROCMIDDAY','LYCANROCMIDDAYFORM','LYCANROCFORMADIURNA','LYCANROCCREPUSCULO','LYCANROCDUSK','LYCANROCDUSKFORM','LYCANROCFORMACREPUSCULO']: return 'LYCANROC'
    if s in ['LYCANROCNOTURNO','LYCANROCMIDNIGHT','LYCANROCMIDNIGHTFORM','LYCANROCFORMANOTURNA']: return 'LYCANROCNOTURNO'
    if s in ['ZACIANHERO','ZACIANCOROADO','ZACIANCROWNEDSWORD']: return 'ZACIAN'
    if s in ['ZAMAZENTAHERO','ZAMAZENTACOROADO','ZAMAZENTACROWNEDSHIELD']: return 'ZAMAZENTA'
    return s
regional={'PERRSERKER':'GALAR','SIRFETCHD':'GALAR','CURSOLA':'GALAR','OBSTAGOON':'GALAR','RUNERIGUS':'GALAR','MRRIME':'GALAR','SNEASLER':'HISUI','OVERQWIL':'HISUI','CLODSIRE':'PALDEA'}
def region(s): return regional.get(norm(s)) or next((x for x in ['ALOLA','GALAR','HISUI','PALDEA'] if norm(s).endswith(x)),None)
def signature(m): return '|'.join(sorted(set(map(norm,m))))
table=read('app/src/main/assets/pvp/master_iv_table.json')
cat={signature(r['key'].split('|')):r for r in table}
families=read('app/src/main/assets/pokemon/families.json')
families={k:([v] if isinstance(v,str) else v) for k,v in families.items()}
names=read('app/src/main/assets/pokemon/names.json')
audit=read('.tmp-stadium-audit.json')
site=read('.tmp-stadium-data.json')
siteidx={raw(m['id']):m for m in site}
sitealiases={'GIRATINAORIGEM':'GIRATINAORIGIN','LYCANROCNOTURNO':'LYCANROCMIDNIGHT','SHAYMINCEU':'SHAYMINSKY','DEOXYSDEFESA':'DEOXYSDEFENSE'}
for m in site:
    key=raw(m['id'])
    for suffix,reg in [('ALOLAN','ALOLA'),('GALARIAN','GALAR'),('HISUIAN','HISUI'),('PALDEAN','PALDEA')]:
        if key.endswith(suffix): siteidx.setdefault(key[:-len(suffix)]+reg,m)
s= (ROOT/'.tmp-stadium-app.js').read_text(encoding='utf8')
cpm=[float(v) for v in re.search(r'cpm:\[([^\]]+)\]',s[s.index('7846:'):])[1].split(',')][78]

def proposals(name):
    key=raw(name);m=siteidx.get(sitealiases.get(key,key))
    if not m or not all(m.get('stats',{}).get(k,0)>0 for k in ['attack','defense','stamina']):return ['']*4,'','Sem atributos/correspondência inequívoca no Stadium; validar manualmente.'
    a,d,h=(m['stats'][k] for k in ['attack','defense','stamina'])
    result=[]
    for total in [44,43,42,41]:
        candidates=[(i,j,total-i-j) for i in range(16) for j in range(16) if 0<=total-i-j<=15]
        score=lambda iv:(a+iv[0])*(d+iv[1])*cpm*cpm*math.floor((h+iv[2])*cpm)
        best=max(map(score,candidates))
        result.append(' ou '.join('/'.join(map(str,v)) for v in candidates if abs(score(v)-best)<1e-7))
    return result,'https://www.stadiumgaming.gg/rank-checker?pokemon='+quote(m['id']),'Proposta para o alvo indicado, por produto de atributos no nível 40; confirmar evolução/forma.'

def entries(members,selected):
    sel=norm(selected);mem=list(dict.fromkeys(map(norm,members)))
    if sel in ['LYCANROC','LYCANROCNOTURNO']: return [cat[sel]] if sel in cat else []
    regional_family=any(region(x) for x in members+[selected])
    if regional_family:
        if sel not in mem and sel not in cat:return []
        scoped=[x for x in dict.fromkeys([sel]+mem) if region(x)==region(sel)]
        if sel in cat:return [cat[sel]]
    else:scoped=mem
    keys=[signature(scoped)]+scoped
    if 'DEOXYS' in scoped:keys += [k for k in cat if k.startswith('DEOXYS')]
    if 'PALAFIN' in scoped:keys += ['PALAFINZERO','PALAFINHERO']
    return [cat[k] for k in dict.fromkeys(keys) if k in cat]

wb=Workbook();wb.remove(wb.active)
ws=wb.create_sheet('Orientações')
for row in [
['PLANILHA DE REVISÃO IV MASTER','27/09/2026'],
['Base confirmada','Nível 40; comparar melhor produto de atributos dentro da mesma soma de IVs.'],
['Regra 67%','10/10/10 é intencional. Não é divergência nem pendência nesta planilha.'],
['Uso','Preencher Decisão, IV aprovado e Observações. Nenhuma alteração automática no app.'],
['Divergências','32 combinações em 22 registros, excluindo os 7 casos já explicados pela evolução.'],
['Regras de evolução','7 diferenças de comparação direta: Farfetch’d Galar → Sirfetch’d e Glimmet → Glimmora. Não consideradas erros.'],
['Famílias sem regra','Ramificações de families.json sem regra efetiva para 98/96/93/91, simulando o filtro regional do catálogo.'],
['Alvo sugerido','Último membro da ramificação cadastrada; CONFIRMAR se há evoluções alternativas. Propostas são preliminares.'],
['Formas herdadas','Formas entre parênteses que usam outra regra da família. Não significa ausência efetiva; revisar se precisam de regra própria.'],
['Cadastro fora das famílias','Nomes de names.json sem participação em families.json; faltas de regra direta são separadas, pois a leitura pode ter contexto de doces.'],
['Pendências','Aegislash sem atributos no Stadium; diferenças de atributos e formas registradas para análise.'],
['Fonte','https://www.stadiumgaming.gg/js/app.js'],
['Snapshot SHA-256',audit['bundle_sha256']],
['Limite','Não é auditoria de disponibilidade no jogo, nem ranking de desempenho em batalhas. Empates são preservados.'],
]:ws.append(row)

fmt=lambda iv:'/'.join(map(str,iv))
headers=['ID','Registro','IV %','IV no app','Proposta Stadium nível 40','Perda produto %','Alvo/observação','Fonte','Decisão','IV aprovado','Observações do usuário']
for title,explained in [('Divergências',False),('Regras de evolução',True)]:
    w=wb.create_sheet(title);w.append(headers)
    for r in audit['rows']:
        if r['level']!=40 or r['percent']=='67' or not r['different']:continue
        known=r['key'] in ['GALARIANFARFETCHD','GLIMMET']
        if known!=explained:continue
        note={'GALARIANFARFETCHD':'Regra atual coincide com Sirfetch’d; diferença direta não implica erro.','GLIMMET':'Regra atual coincide com Glimmora; diferença direta não implica erro.'}.get(r['key'],'Confirmar espécie/forma alvo da regra familiar antes de corrigir.')
        w.append([f"{'E' if explained else 'D'}-{w.max_row:03}",r['key'],int(r['percent']),fmt(r['app']),' ou '.join(map(fmt,r['best'])),round(r['loss'],6),note,'https://www.stadiumgaming.gg/rank-checker?pokemon='+quote(r['site']),'Pendente','',''])

w=wb.create_sheet('Famílias sem regra');w.append(['ID','Família/ramificação','Membros afetados','Alvo sugerido (confirmar)','Percentuais ausentes','Proposta 98','Proposta 96','Proposta 93','Proposta 91','Fonte','Nota','Decisão','IV 98 aprovado','IV 96 aprovado','IV 93 aprovado','IV 91 aprovado','Observações do usuário'])
seen=set();inherited=[]
for members in dict.fromkeys(tuple(v) for v in families.values()):
    members=list(members)
    for reg in dict.fromkeys(region(x) for x in members):
        branch=[x for x in members if region(x)==reg]
        missing={}
        for name in branch:
            rules=entries(members,name)
            gaps=tuple(p for p in ['98','96','93','91'] if not any(p in rule for rule in rules))
            if gaps:missing.setdefault(gaps,[]).append(name)
            elif '(' in name and norm(name) not in cat and not norm(name).startswith('LYCANROC'):
                inherited.append((name,branch,', '.join(r['key'] for r in rules)))
        for gaps,affected in missing.items():
            sig=(signature(branch),gaps,tuple(affected))
            if sig in seen:continue
            seen.add(sig);target=branch[-1]
            vals,url,note=proposals(target)
            w.append([f'F-{w.max_row:03}',' / '.join(branch),'; '.join(affected),target,', '.join(gaps),*vals,url,note,'Pendente','','','','',''])

w=wb.create_sheet('Formas herdadas');w.append(['ID','Forma','Família/ramificação','Regras utilizadas','Nota','Decisão','Observações do usuário'])
for name,branch,rules in dict.fromkeys((n,tuple(b),r) for n,b,r in inherited):
    w.append([f'H-{w.max_row:03}',name,' / '.join(branch),rules,'Já possui cobertura familiar; conferir se os atributos desta forma exigem combinações próprias.','Pendente',''])

w=wb.create_sheet('Fora das famílias');w.append(['ID','Espécie/forma','National Dex','Nota','Decisão','Observações do usuário'])
family_names={raw(x) for v in families.values() for x in v}|{raw(k) for k in families}
for row in names:
    name=row['name']
    if raw(name) not in family_names and not entries([name],name):
        w.append([f'N-{w.max_row:03}',name,row['dex'],'Sem associação explícita neste levantamento e sem regra direta. Contexto de doces pode fornecer família; confirmar cadastro.','Pendente',''])

w=wb.create_sheet('Pendências de dados');w.append(['Item','Situação','Decisão','Observações do usuário'])
for row in [['Aegislash','Stadium não fornece atributos no registro genérico. Confirmar Escudo/Espada antes de propor IVs.'],['Dugtrio','Defesa 134 no registro genérico do Stadium versus 136 no app. Não substituir atributos auditados sem confirmar a fonte.'],['Gourgeist','Stadium genérico 171/219/146 versus app 175/213/163. Confirmar tamanho/forma.']]:w.append(row+['Pendente',''])

for number,w in enumerate(wb.worksheets):
    w.freeze_panes='A2';w.sheet_view.showGridLines=False
    for cell in w[1]:cell.fill=PatternFill('solid',fgColor='243B53');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True,vertical='center')
    w.row_dimensions[1].height=34
    for col in w.columns:
        label=col[0].value or ''
        width=18
        if any(x in label for x in ['Nota','Observa','Família','Membros','Situação','Alvo/']):width=55
        elif any(x in label for x in ['Registro','Espécie','Forma','Alvo sugerido','Regras']):width=30
        elif label=='Fonte':width=25
        w.column_dimensions[col[0].column_letter].width=width
    if number==0:
        w.column_dimensions['A'].width=30;w.column_dimensions['B'].width=110
    elif w.max_row>1:
        t=Table(displayName=f'Revisao{number}',ref=w.dimensions);t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);w.add_table(t)
    for row in w.iter_rows(min_row=2):
        w.row_dimensions[row[0].row].height=62 if number else 45
        for c in row:
            c.alignment=Alignment(wrap_text=True,vertical='top')
            if isinstance(c.value,str) and c.value.startswith('https://'):
                c.hyperlink=c.value;c.font=Font(color='0563C1',underline='single')
    for c in w[1]:
        if c.value=='Decisão':
            dv=DataValidation(type='list',formula1='"Pendente,Aprovar proposta,Manter atual,Definir outro IV,Revisar forma/família"');w.add_data_validation(dv);dv.add(f'{c.column_letter}2:{c.column_letter}{max(2,w.max_row)}')
        if c.value in ['Decisão','IV aprovado','Observações do usuário'] or str(c.value).endswith('aprovado'):
            for cells in w.iter_rows(min_row=2,min_col=c.column,max_col=c.column):cells[0].fill=PatternFill('solid',fgColor='FFF2CC')

path=ROOT/'docs/revisao-iv-master-nivel-40.xlsx';wb.save(path)
check=load_workbook(path)
assert check['Divergências'].max_row-1==32
assert not any('Kangaskhan' in str(r[2]) for r in check['Famílias sem regra'].iter_rows(min_row=2,values_only=True))
assert all(len(str(r[3]))>1 for r in check['Famílias sem regra'].iter_rows(min_row=2,values_only=True))
assert check['Regras de evolução'].max_row-1==7
assert any('Growlithe (Hisui)' in str(r[2]) for r in check['Famílias sem regra'].iter_rows(min_row=2,values_only=True))
print(path)
for w in check:print(w.title,w.max_row-1)
